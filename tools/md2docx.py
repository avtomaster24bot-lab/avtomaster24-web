import re, sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
src, out = sys.argv[1], sys.argv[2]
lines = open(src, encoding='utf-8').read().split('\n')
d = Document()
st = d.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
for s in d.sections: s.left_margin = s.right_margin = Cm(2); s.top_margin = s.bottom_margin = Cm(1.8)
def runs(p, text):
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)', text):
        if not part: continue
        if part.startswith('**'): r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('`'): r = p.add_run(part[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(9.5)
        elif part.startswith('*') and len(part) > 2: r = p.add_run(part[1:-1]); r.italic = True
        else: p.add_run(part)
i = 0
while i < len(lines):
    l = lines[i]
    if l.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            if not re.match(r'^\|[-| ]+\|$', lines[i]): rows.append([c.strip() for c in lines[i].strip('|').split('|')])
            i += 1
        t = d.add_table(rows=len(rows), cols=len(rows[0])); t.style = 'Light Grid Accent 1'
        for ri, r in enumerate(rows):
            for ci, c in enumerate(r):
                cell = t.cell(ri, ci); cell.text = ''; p = cell.paragraphs[0]; runs(p, c)
                for rr in p.runs: rr.font.size = Pt(9); rr.bold = rr.bold or ri == 0
        d.add_paragraph(); continue
    if l.startswith('# '): h = d.add_heading(l[2:], 0)
    elif l.startswith('## '): d.add_heading(l[3:], 1)
    elif l.startswith('### '): d.add_heading(l[4:], 2)
    elif l.startswith('> '): p = d.add_paragraph(); runs(p, l[2:]); p.paragraph_format.left_indent = Cm(0.6); [setattr(r.font.color, 'rgb', RGBColor(0x7a, 0x55, 0x00)) for r in p.runs]
    elif re.match(r'^\s*- ', l):
        lvl = (len(l) - len(l.lstrip())) // 2; p = d.add_paragraph(style='List Bullet 2' if lvl else 'List Bullet'); runs(p, l.strip()[2:])
    elif re.match(r'^\d+\. ', l): p = d.add_paragraph(style='List Number'); runs(p, re.sub(r'^\d+\. ', '', l))
    elif l.strip() == '---': pass
    elif l.strip(): p = d.add_paragraph(); runs(p, l)
    i += 1
d.save(out); print('saved', out)
